"""
app/services/parser.py

文件解析服務：
- parse_docx() → python-docx，保留標題層級 + 表格 Markdown

⚠ parse_pdf_text()（pdfplumber）已移除。
  PDF 現在直接送 Gemini Flash 視覺理解，不再走文字提取層。
  詳見 gemini.convert_pdf_to_chunks()。
"""
import io
from typing import Any

from app.core.logging import get_logger

logger = get_logger("parser")


def _table_to_markdown(rows: list[list[Any]]) -> str:
    """二維陣列 → Markdown table 字串。"""
    if not rows:
        return ""
    lines = []
    for i, row in enumerate(rows):
        cells = [str(cell or "").strip() for cell in row]
        lines.append("| " + " | ".join(cells) + " |")
        if i == 0:
            lines.append("| " + " | ".join(["---"] * len(cells)) + " |")
    return "\n".join(lines)


def parse_docx(file_bytes: bytes) -> dict[str, Any]:
    """
    python-docx 解析器。
    保留 Heading 1/2/3 層級，表格轉 Markdown。

    Returns:
        {
            "headings": [{"level": 1, "text": "..."}, ...],
            "paragraphs": ["...", ...],
            "tables": ["| col1 | col2 |\n...", ...],
            "full_text": "完整純文字（供 Gemini 處理）"
        }
    """
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(io.BytesIO(file_bytes))

    headings: list[dict] = []
    paragraphs: list[str] = []
    tables: list[str] = []
    full_text_parts: list[str] = []

    for element in doc.element.body:
        tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag

        if tag == "p":
            para_text = "".join(r.text for r in element.findall(f".//{qn('w:t')}"))
            if not para_text.strip():
                continue

            style_name = ""
            style_elem = element.find(f".//{qn('w:pStyle')}")
            if style_elem is not None:
                style_name = style_elem.get(qn("w:val"), "")

            if style_name.startswith("Heading"):
                try:
                    level = int(style_name.replace("Heading", "").strip())
                except ValueError:
                    level = 1
                if level <= 3:
                    headings.append({"level": level, "text": para_text.strip()})
                    full_text_parts.append(f"{'#' * level} {para_text.strip()}")
                    continue

            paragraphs.append(para_text.strip())
            full_text_parts.append(para_text.strip())

        elif tag == "tbl":
            pass  # 下方用 doc.tables 統一處理

    for table in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        md = _table_to_markdown(rows)
        if md:
            tables.append(md)
            full_text_parts.append(md)

    result = {
        "headings": headings,
        "paragraphs": paragraphs,
        "tables": tables,
        "full_text": "\n\n".join(full_text_parts),
    }
    logger.info(
        f"python-docx 解析完成：{len(headings)} 標題, {len(paragraphs)} 段落, {len(tables)} 表格"
    )
    return result
